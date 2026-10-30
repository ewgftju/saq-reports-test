#!/usr/bin/env python3
"""Compare actual report exports with the private customer templates, read-only.

Run from any directory:
  python scripts/verify-source-layout.py --received /path/to/upload \
    --professional /path/to/unpacked --output /path/to/scratch/exports

The existing application generates XLSX files. openpyxl/xlrd only read them.
No source observation, BIN, person, address or formula is written to the repo.
Only structural allowlisted coordinates are compared; failures print coordinates,
never private cell values. Dates/periods are intentionally parameterized.
"""
from __future__ import annotations

import argparse
import json
import math
import os
from pathlib import Path
import re
import struct
import subprocess
import sys

import openpyxl
import xlrd
from openpyxl.utils.cell import get_column_letter, range_boundaries


# id, file, meaningful last column, complete static prefix, fixed body range,
# body label columns, parameterized cells. No source observations are included.
RECEIVED = [
    ('received-plan-detail', 'Информация исполнению перечня отчеты ДВГА.xls', 14, 8, None, (), {'A4'}),
    ('received-plan-monitoring', 'Информация по исполнению Перечня.xlsx', 28, 11, (12, 37), (2, 3), {'B4', 'B6', 'D8'}),
    ('received-plan-summary', 'Информация сводная по Перечню.xlsx', 15, 6, (7, 28), (2, 3), {'D4'}),
    ('received-assignment-decisions', 'Отчет по соггласованным поручениям.xlsx', 6, 11, (12, 32), (1, 2), set()),
    ('received-electronic-rating', 'Применение ЭВГА.xlsx', 9, 9, (10, 32), (2, 3), {'B5'}),
    ('received-quality-journal', 'Журнал ДВГА- несоотв на 03.07.2026 (1).xlsx', 12, 4, None, (), {'F1'}),
]
PROFESSIONAL = [
    ('prof-halfyear-list', 'EGOV полугод список каз рус.xlsx', 39, 1, None, (), set()),
    ('prof-list-execution', 'Актуальная инф по перечню.xlsx', 22, 2, (3, 23), (1, 2), {'C1'}),
    ('prof-audit-organizations', 'Анализ АО.xlsx', 38, 4, None, (), set()),
    ('prof-public-interest', 'Анализ ОПИ 1-2025.xlsx', 46, 5, None, (), set()),
    ('prof-administrative', 'информация по ст КоАП.xlsx', 12, 3, None, (), set()),
    ('prof-year-comparison', 'Таблица Упр Анализа.xlsx', 9, 2, (3, 19), (2, 3, 4, 5), {'F1', 'G1'}),
]


def canonical(value):
    if value is None:
        return ''
    if isinstance(value, str):
        return value.replace('\r\n', '\n').replace('\r', '\n').strip()
    return value


class Source:
    def __init__(self, path):
        self.path = path
        self.xls = path.suffix.lower() == '.xls'
        if self.xls:
            self.book = xlrd.open_workbook(path, formatting_info=True)
            self.sheet = self.book.sheet_by_index(0)
            self.names = self.book.sheet_names()
            self.max_column = self.sheet.ncols
            self.merges = {f'{get_column_letter(c+1)}{r+1}:{get_column_letter(cc)}{rr}'
                           for r, rr, c, cc in self.sheet.merged_cells}
        else:
            self.book = openpyxl.load_workbook(path, data_only=False)
            self.sheet = self.book.worksheets[0]
            self.names = self.book.sheetnames
            self.max_column = self.sheet.max_column
            self.merges = {str(m) for m in self.sheet.merged_cells.ranges}

    def value(self, row, col):
        if self.xls:
            return self.sheet.cell_value(row-1, col-1) if row <= self.sheet.nrows and col <= self.sheet.ncols else None
        return self.sheet.cell(row, col).value

    def explicit_width(self, col):
        if self.xls:
            info = self.sheet.colinfo_map.get(col-1)
            return info.width/256 if info else None
        for dim in self.sheet.column_dimensions.values():
            if dim.min <= col <= dim.max:
                return dim.width
        return None

    def row_height(self, row):
        if self.xls:
            info = self.sheet.rowinfo_map.get(row-1)
            return (info.height if info else self.sheet.default_row_height)/20
        dim = self.sheet.row_dimensions.get(row)
        return (dim.height if dim and dim.height is not None else self.sheet.sheet_format.defaultRowHeight)

    def style(self, row, col):
        if self.xls:
            xf = self.book.xf_list[self.sheet.cell_xf_index(row-1, col-1)]
            font, align = self.book.font_list[xf.font_index], xf.alignment
            sides = {0:None,1:'thin',2:'medium',3:'dashed',4:'dotted',5:'thick',6:'double',7:'hair',8:'mediumDashed',9:'dashDot',10:'mediumDashDot',11:'dashDotDot',12:'mediumDashDotDot',13:'slantDashDot'}
            borders=[sides[getattr(xf.border,edge+'_line_style')] for edge in ('left','right','top','bottom')]
            # An Excel merged anchor displays the outer right/bottom borders
            # from the final cell, not only the anchor's own BIFF XF record.
            for ref in self.merges:
                c1,r1,c2,r2=range_boundaries(ref)
                if (row,col)==(r1,c1):
                    end=self.book.xf_list[self.sheet.cell_xf_index(r2-1,c2-1)].border
                    borders[1]=borders[1] or sides[end.right_line_style]
                    borders[3]=borders[3] or sides[end.bottom_line_style]
            return {'size':font.height/20,'bold':bool(font.bold),'italic':bool(font.italic),
                    'horizontal':{0:'general',1:'left',2:'center',3:'right',4:'fill',5:'justify',6:'centerContinuous',7:'distributed'}[align.hor_align],
                    'vertical':{0:'top',1:'center',2:'bottom',3:'justify',4:'distributed'}[align.vert_align],
                    'wrap':bool(align.text_wrapped),'rotation':align.rotation,'indent':align.indent_level,
                    'borders':tuple(borders)}
        return cell_style(self.sheet.cell(row,col))

    def page(self):
        if not self.xls:
            setup={key:getattr(self.sheet.page_setup,key) for key in PAGE_FIELDS if getattr(self.sheet.page_setup,key) is not None}
            margins={key:getattr(self.sheet.page_margins,key) for key in MARGIN_FIELDS}
            return setup,margins
        # Read the source BIFF print records independently of source-layout.py.
        from xlrd.compdoc import CompDoc
        compound=CompDoc(self.path.read_bytes())
        data=compound.get_named_stream('Workbook') or compound.get_named_stream('Book')
        pos=self.book._sh_abs_posn[0]-self.book.base
        setup,margins={},{}
        while pos+4<=len(data):
            opcode,length=struct.unpack_from('<HH',data,pos);chunk=data[pos+4:pos+4+length];pos+=length+4
            if opcode==0x000A:break
            if opcode in (0x26,0x27,0x28,0x29):margins[('left','right','top','bottom')[opcode-0x26]]=struct.unpack_from('<d',chunk)[0]
            if opcode==0xA1 and length>=34:
                paper,scale,first,width,height,flags,hdpi,vdpi,header,footer,_=struct.unpack_from('<8H2dH',chunk)
                setup={'paperSize':paper,'scale':scale,'firstPageNumber':first,'fitToWidth':width,'fitToHeight':height,
                       'orientation':'portrait' if flags&2 else 'landscape','blackAndWhite':bool(flags&8),'draft':bool(flags&16),
                       'useFirstPageNumber':bool(flags&128),'pageOrder':'overThenDown' if flags&1 else 'downThenOver',
                       'horizontalDpi':hdpi,'verticalDpi':vdpi}
                margins.update(header=header,footer=footer)
        return setup,margins


PAGE_FIELDS=('orientation','paperSize','scale','fitToWidth','fitToHeight','firstPageNumber','useFirstPageNumber','blackAndWhite','draft','horizontalDpi','verticalDpi','pageOrder')
MARGIN_FIELDS=('left','right','top','bottom','header','footer')


def cell_style(cell):
    a=cell.alignment
    return {'size':cell.font.sz,'bold':bool(cell.font.b),'italic':bool(cell.font.i),
            'horizontal':a.horizontal or 'general','vertical':a.vertical or 'bottom','wrap':bool(a.wrap_text),
            'rotation':a.text_rotation or 0,'indent':a.indent or 0,
            'borders':tuple(getattr(cell.border,edge).style if getattr(cell.border,edge) else None for edge in ('left','right','top','bottom'))}


def output_width(sheet, col):
    for dim in sheet.column_dimensions.values():
        if dim.min <= col <= dim.max:
            return dim.width
    return sheet.sheet_format.defaultColWidth


def is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def compare_one(spec, source_dir, export_dir):
    ident, filename, width, prefix, fixed, label_cols, dynamic = spec
    source = Source(source_dir / filename)
    exported = openpyxl.load_workbook(export_dir / f'{ident}.xlsx', data_only=False)
    failures = []
    check = lambda condition, message: failures.append(message) if not condition else None
    check(exported.sheetnames == source.names, 'sheet names/count differ from source')
    check(len(exported.worksheets) == 1, 'extra worksheet')
    sheet = exported.worksheets[0]
    # Some sources contain only empty formatted cells beyond the table. Keeping
    # or trimming those is permitted, adding/deleting meaningful fields is not.
    check(width <= sheet.max_column <= source.max_column, f'width outside source boundaries: {width}..{source.max_column}, got {sheet.max_column}')
    for row in sheet.iter_rows(min_col=width+1):
        check(all(canonical(cell.value)=='' for cell in row), 'extra meaningful field beyond source table')
    static_cells=[]
    for row in range(1, prefix+1):
        for col in range(1, width+1):
            ref = f'{get_column_letter(col)}{row}'
            expected, actual = source.value(row, col), sheet.cell(row, col).value
            if ref in dynamic:
                selected_year='2025' if ident=='prof-year-comparison' and ref=='G1' else '2026'
                check(actual is not None and selected_year in str(actual), f'{ref}: selected reporting year is missing')
                continue
            check(canonical(expected) == canonical(actual), f'{ref}: static prefix/title/numbering differs')
            if canonical(expected)!='':static_cells.append((row,col))
    if fixed:
        for row in range(fixed[0], fixed[1]+1):
            for col in label_cols:
                ref = f'{get_column_letter(col)}{row}'
                check(canonical(source.value(row, col)) == canonical(sheet.cell(row, col).value), f'{ref}: fixed row label/position differs')
                if canonical(source.value(row,col))!='':static_cells.append((row,col))
    last_static = fixed[1] if fixed else prefix
    expected_merges = {ref for ref in source.merges if range_boundaries(ref)[3] <= last_static}
    actual_merges = {str(m) for m in sheet.merged_cells.ranges if m.max_row <= last_static}
    check(expected_merges == actual_merges, f'static merges differ; missing={sorted(expected_merges-actual_merges)}, extra={sorted(actual_merges-expected_merges)}')
    for col in range(1, width+1):
        expected = source.explicit_width(col)
        if expected is not None:
            actual = output_width(sheet, col)
            check(actual is not None and abs(expected-actual) < .01, f'{get_column_letter(col)}: original column width differs')
    for row in range(1,last_static+1):
        expected=source.row_height(row)
        actual=sheet.row_dimensions[row].height or sheet.sheet_format.defaultRowHeight
        if expected is not None:check(actual is not None and abs(actual-expected)<.01,f'row {row}: original height differs')
    for row,col in static_cells:
        expected,actual=source.style(row,col),cell_style(sheet.cell(row,col))
        changed=[key for key in expected if expected[key]!=actual[key]]
        check(not changed,f'{get_column_letter(col)}{row}: source style differs ({", ".join(changed)})')
    setup,margins=source.page()
    for key,expected in setup.items():check(getattr(sheet.page_setup,key)==expected,f'pageSetup.{key}: source setting differs')
    for key,expected in margins.items():
        actual=getattr(sheet.page_margins,key)
        check(actual is not None and abs(actual-expected)<.000001,f'pageMargins.{key}: source margin differs')
    for row in sheet:
        for cell in row:
            if cell.value is not None:
                check(cell.font.name == 'Times New Roman', f'{cell.coordinate}: font is not Times New Roman')
                check(cell.data_type != 'e', f'{cell.coordinate}: Excel error value')
                check(not (cell.data_type == 'f' and re.search(r'#REF!|\[', str(cell.value))), f'{cell.coordinate}: broken/external formula')
    # Source totals may move with variable-length detail, but keep their labels
    # and their columns; technical reference sheets/totals must not be added.
    if ident == 'prof-administrative':
        check(not any(isinstance(c.value, str) and re.fullmatch(r'\s*(итого|всего)\s*[:.]?\s*', c.value, re.I)
                      for row in sheet.iter_rows(min_row=prefix+1) for c in row), 'extra total row in administrative report')
    for total_ref in {'prof-audit-organizations': ['A7'], 'prof-public-interest': ['C40', 'C42'],
                      'received-plan-detail': ['B12']}.get(ident, []):
        col_letters, source_row = re.fullmatch(r'([A-Z]+)(\d+)', total_ref).groups()
        source_col = openpyxl.utils.cell.column_index_from_string(col_letters)
        label = canonical(source.value(int(source_row), source_col))
        check(any(canonical(sheet.cell(r, source_col).value) == label for r in range(prefix+1, sheet.max_row+1)),
              f'{total_ref}: source total label missing or moved to another column')
    if ident == 'received-plan-monitoring':
        check_monitor_totals(sheet, check)
    return {'form': ident, 'failures': list(dict.fromkeys(failures))}


def check_monitor_totals(sheet, check):
    counts = ['D','E','F','G','H','I','J','K','L','M','N','O','P','Q','R','S','T','U','V','W','Y','AA']
    for target, children in [(32, range(12,32)), (36, range(33,36)), (37, [32,36])]:
        for col in counts:
            values = [sheet[f'{col}{r}'].value for r in children]
            total = sheet[f'{col}{target}'].value
            nonempty=[v for v in values if canonical(v)!='']
            if nonempty and all(is_number(v) for v in nonempty):
                check(is_number(total) and abs(total-sum(nonempty)) < .00001, f'{col}{target}: subtotal does not reconcile')
            elif nonempty:
                check(not is_number(total), f'{col}{target}: incomplete population shown as numeric total')
            else:
                check(canonical(total)=='' or total==0, f'{col}{target}: empty template slots have a nonzero total')
        for percentage, part in [('X','W'), ('Z','Y'), ('AB','AA')]:
            numerator, denominator = sheet[f'{part}{target}'].value, sheet[f'V{target}'].value
            actual = sheet[f'{percentage}{target}'].value
            if is_number(numerator) and is_number(denominator) and denominator > 0:
                scale=1 if re.search(r'(?<!\\)%',sheet[f'{percentage}{target}'].number_format) else 100
                check(is_number(actual) and abs(actual-numerator/denominator*scale) < .011/100*scale,
                      f'{percentage}{target}: percentage is not weighted from total counts')
            else:
                check(not is_number(actual), f'{percentage}{target}: undefined percentage presented as a number')


def generate_exports(repo, output):
    output.mkdir(parents=True, exist_ok=True)
    script = """
import {writeFile} from 'node:fs/promises';
import {archiveForms} from './form-registry.js';
import {generateArchiveReport} from './archive-engine.js';
import {exportArchiveXlsx} from './export.js';
for(const form of archiveForms.filter(f=>['received','professional'].includes(f.engine)))
 await writeFile(`${process.argv[1]}/${form.id}.xlsx`,exportArchiveXlsx(generateArchiveReport(form.id)));
"""
    subprocess.run([os.environ.get('CODEX_PRIMARY_RUNTIME_NODE', 'node'), '--input-type=module', '-e', script, str(output)], cwd=repo, check=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--received', type=Path, required=True)
    parser.add_argument('--professional', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--no-generate', action='store_true')
    args = parser.parse_args()
    repo = Path(__file__).resolve().parents[1]
    output = args.output.resolve()
    if not args.no_generate:
        generate_exports(repo, output)
    results = [compare_one(spec, folder.resolve(), output) for specs, folder in [(RECEIVED,args.received), (PROFESSIONAL,args.professional)] for spec in specs]
    print(json.dumps({'checked':len(results), 'passed':sum(not r['failures'] for r in results), 'results':results,
                      'scope':'Twelve supplied report templates. The separate KoAP reference workbook is not an extra report worksheet. Earlier UAO originals are outside this audit.'}, ensure_ascii=False, indent=2))
    return 1 if any(r['failures'] for r in results) else 0


if __name__ == '__main__':
    sys.exit(main())
