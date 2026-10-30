#!/usr/bin/env python3
"""Extract only allowlisted headers of the six report templates received 30.10.2026.

Usage: python scripts/extract-received-forms.py /path/to/upload [received-schemas.js]
No observation rows, names, BINs, customer dates or comments are copied.
openpyxl and xlrd are used exclusively for reading supplied workbooks.
"""
import argparse
import json
import re
from importlib.machinery import SourceFileLoader
from pathlib import Path
import openpyxl
import xlrd
from openpyxl.utils.cell import get_column_letter, range_boundaries
extract_source_layout = SourceFileLoader('source_layout', str(Path(__file__).with_name('source-layout.py'))).load_module().extract_source_layout

MANIFEST = [
    ('Информация исполнению перечня отчеты ДВГА.xls', 'received-plan-detail', 'Исполнение Перечня объектов государственного аудита — отчет ДВГА', 'digital-risks', 2, 6, 7, 14, 'Исполнение Перечня'),
    ('Информация по исполнению Перечня.xlsx', 'received-plan-monitoring', 'Мониторинг исполнения Перечня объектов государственного аудита', 'digital-risks', 2, 8, 10, 28, 'Мониторинг Перечня'),
    ('Информация сводная по Перечню.xlsx', 'received-plan-summary', 'Сводная информация по Перечню объектов государственного аудита', 'digital-risks', 2, 4, 6, 15, 'Сводная по Перечню'),
    ('Отчет по соггласованным поручениям.xlsx', 'received-assignment-decisions', 'Отчет по согласованным поручениям', 'digital-risks', 3, 9, 10, 6, 'Согласованные поручения'),
    ('Применение ЭВГА.xlsx', 'received-electronic-rating', 'Рейтинг применения электронного аудита', 'digital-risks', 5, 7, 8, 9, 'Применение ЭВГА'),
    ('Журнал ДВГА- несоотв на 03.07.2026 (1).xlsx', 'received-quality-journal', 'Журнал материалов, не соответствующих Стандартам ГАФК', 'quality', 1, 2, 4, 12, 'Несоответствия Стандартам'),
]


def clean(value):
    if isinstance(value, (int, float)):
        return int(value) if value == int(value) else value
    return re.sub(r'\s+', ' ', str(value)).strip()


def extract(path, first, last, column_count):
    if path.suffix == '.xls':
        workbook = xlrd.open_workbook(path, formatting_info=True)
        source = workbook.sheet_by_index(0)
        source_name = source.name
        values = {(r + 1, c + 1): source.cell_value(r, c) for r in range(first - 1, last) for c in range(min(column_count, source.ncols))}
        source_merges = [f'{get_column_letter(c + 1)}{r + 1}:{get_column_letter(cc)}{rr}' for r, rr, c, cc in source.merged_cells]
    else:
        workbook = openpyxl.load_workbook(path, data_only=False)
        source = workbook.worksheets[0]
        source_name = source.title
        values = {(r, c): source.cell(r, c).value for r in range(first, last + 1) for c in range(1, column_count + 1)}
        source_merges = [str(item) for item in source.merged_cells.ranges]
    header = {}
    for (row, column), value in values.items():
        if value is None or value == '' or isinstance(value, str) and value.startswith('='):
            continue
        if not isinstance(value, str) or value.strip():
            header[f'{get_column_letter(column)}{row - first + 1}'] = value
    merges = []
    for ref in source_merges:
        c1, r1, c2, r2 = range_boundaries(ref)
        if first <= r1 <= r2 <= last and c2 <= column_count:
            merges.append(f'{get_column_letter(c1)}{r1 - first + 1}:{get_column_letter(c2)}{r2 - first + 1}')
    columns = []
    for col in range(1, column_count + 1):
        labels = []
        for row in range(1, last - first + 2):
            value = header.get(f'{get_column_letter(col)}{row}')
            if value is None:
                for merge in merges:
                    c1, r1, c2, r2 = range_boundaries(merge)
                    if c1 <= col <= c2 and r1 <= row <= r2:
                        value = header.get(f'{get_column_letter(c1)}{r1}')
                        break
            if value is not None and str(value) not in labels:
                labels.append(str(value))
        columns.append({'letter': get_column_letter(col), 'labels': labels})
    return {'sourceSheetName': source_name, 'kind': 'table', 'columnCount': column_count, 'headerDepth': last - first + 1, 'headerCells': header, 'merges': sorted(merges), 'columns': columns, 'sourceHeaderRows': [first, last]}


LAYOUTS = {
    'received-plan-detail': {'prefixEnd': 8, 'rowCount': 13, 'dataStart': 9, 'dataEnd': 11, 'totalRow': 12, 'labelColumns': ['B'], 'fixedRows': [12]},
    'received-plan-monitoring': {'prefixEnd': 11, 'rowCount': 47, 'labelColumns': ['B', 'C'], 'fixedRows': list(range(12, 38)), 'bodyStart': 12, 'bodyEnd': 37},
    'received-plan-summary': {'prefixEnd': 6, 'rowCount': 28, 'labelColumns': ['B', 'C'], 'fixedRows': list(range(7, 29)), 'bodyStart': 7, 'bodyEnd': 28},
    'received-assignment-decisions': {'prefixEnd': 11, 'rowCount': 32, 'labelColumns': ['A', 'B'], 'fixedRows': list(range(12, 33)), 'bodyStart': 12, 'bodyEnd': 32},
    'received-electronic-rating': {'prefixEnd': 9, 'rowCount': 32, 'labelColumns': ['B', 'C'], 'fixedRows': list(range(10, 33)), 'bodyStart': 10, 'bodyEnd': 32},
    'received-quality-journal': {'prefixEnd': 4, 'rowCount': 5, 'dataStart': 5, 'dataEnd': 5, 'labelColumns': [], 'fixedRows': []},
}


def source_layout(path, form_id, column_count):
    definition = LAYOUTS[form_id]
    allowed = [f'{get_column_letter(column)}{row}' for row in range(1, definition['prefixEnd'] + 1) for column in range(1, column_count + 1)]
    allowed += [f'{column}{row}' for row in definition['fixedRows'] for column in definition['labelColumns']]
    layout = extract_source_layout(path, column_count=column_count, row_count=definition['rowCount'], allowed_cells=allowed)
    # Keep intentional line breaks/spaces in labels. Blank helper cells outside
    # the original print/table rectangle are not report columns.
    return layout, {key: value for key, value in definition.items() if key not in ['rowCount', 'fixedRows']}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('source', type=Path)
    parser.add_argument('output', nargs='?', type=Path, default=Path('received-schemas.js'))
    args = parser.parse_args()
    forms = []
    for filename, form_id, title, department, number, first, last, column_count, sheet_name in MANIFEST:
        sheet = extract(args.source / filename, first, last, column_count)
        layout, placement = source_layout(args.source / filename, form_id, column_count)
        forms.append({'id': form_id, 'name': title, 'sourceFile': filename, 'departmentId': department, 'catalogNumber': number, 'sources': ['ВГА'], 'engine': 'received', 'sheets': [{'id': form_id + '-s1', 'name': sheet['sourceSheetName'], **sheet, 'templateLayout': layout, 'sourcePlacement': placement}]})
    args.output.write_text('// Structural template layout only; supplied customer observation rows are never published.\nexport const receivedForms = ' + json.dumps(forms, ensure_ascii=False, separators=(',', ':')) + ';\n')


if __name__ == '__main__':
    main()
