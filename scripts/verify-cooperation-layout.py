#!/usr/bin/env python3
"""Audit an actual cooperation XLSX export against the private source workbook.

The application writes the scratch export; openpyxl only reads both workbooks.
Source observations and notes are never copied into this repository or printed.
Usage: python scripts/verify-cooperation-layout.py --source /private/source.xlsx \
       --output /scratch/cooperation.xlsx
"""
from __future__ import annotations

import argparse
import json
import math
import os
from pathlib import Path
import re
import subprocess

import openpyxl
from openpyxl.utils.cell import get_column_letter


# Original table boundaries include intentionally blank columns between fields.
# Physical formatting extends farther on two sheets; it is preserved as well.
SHEETS = {
    'ЦГО': {'header': 3, 'total': 34, 'width': 47,
            'dynamic': {'B1', 'I1', 'C2', 'AR2'}},
    'МИО': {'header': 3, 'total': 24, 'width': 51,
            'dynamic': {'L1', 'C2', 'I2', 'M2', 'AV2'}},
    'ДП': {'header': 5, 'total': 26, 'width': 49,
           'dynamic': {'B3', 'D3', 'E3', 'C4', 'D4', 'AT4'}},
    'ИТОГО': {'header': 3, 'total': 7, 'width': 44,
              'dynamic': {'B1', 'C2', 'AO2'}},
}
PRIVATE_NOTES = {'МИО': ['B28'], 'ДП': ['C28', 'C29']}
PAGE_FIELDS = ('orientation', 'paperSize', 'scale', 'fitToWidth', 'fitToHeight',
               'firstPageNumber', 'useFirstPageNumber', 'blackAndWhite', 'draft',
               'horizontalDpi', 'verticalDpi', 'pageOrder')
MARGIN_FIELDS = ('left', 'right', 'top', 'bottom', 'header', 'footer')


def canonical(value):
    return value.replace('\r\n', '\n').replace('\r', '\n').strip() if isinstance(value, str) else value


def style(cell):
    alignment = cell.alignment
    return {
        'size': cell.font.sz, 'bold': bool(cell.font.b), 'italic': bool(cell.font.i),
        'horizontal': alignment.horizontal or 'general',
        'vertical': alignment.vertical or 'bottom', 'wrap': bool(alignment.wrap_text),
        'rotation': alignment.text_rotation or 0, 'indent': alignment.indent or 0,
        'borders': tuple(getattr(cell.border, side).style if getattr(cell.border, side) else None
                         for side in ('left', 'right', 'top', 'bottom')),
        'numberFormat': cell.number_format,
    }


def width(sheet, column):
    for dimension in sheet.column_dimensions.values():
        if dimension.min <= column <= dimension.max:
            return dimension.width
    return sheet.sheet_format.defaultColWidth


def height(sheet, row):
    dimension = sheet.row_dimensions.get(row)
    return dimension.height if dimension and dimension.height is not None else sheet.sheet_format.defaultRowHeight


def expected_label(source, cached, sheet_name, cell):
    value = source[sheet_name][cell].value
    if source[sheet_name][cell].data_type != 'f':
        return value
    if sheet_name == 'МИО' and cell == 'AU2':
        return source['ЦГО']['AQ2'].value
    # Two institution names are formulas linked to unavailable external files;
    # their cached strings are the actual labels displayed in the supplied file.
    if sheet_name == 'ЦГО' and cell in ('B31', 'B32'):
        cached_value = cached[sheet_name][cell].value
        if not isinstance(cached_value, str):
            raise ValueError('Original cached institution label unavailable')
        return cached_value
    raise ValueError(f'Unexpected formula in structural cell {sheet_name}!{cell}')


def compare_sheet(source, cached, exported, name):
    original, actual, spec = source[name], exported[name], SHEETS[name]
    failures = []
    def check(condition, message):
        if not condition:
            failures.append(message)

    check(actual.max_row == original.max_row, 'original row extent differs')
    check(actual.max_column == original.max_column, 'original column extent differs')
    check({str(merge) for merge in actual.merged_cells.ranges} ==
          {str(merge) for merge in original.merged_cells.ranges}, 'original merged ranges differ')
    for column in range(1, original.max_column + 1):
        expected_width, actual_width = width(original, column), width(actual, column)
        if expected_width is not None:
            check(actual_width is not None and abs(expected_width - actual_width) < .01,
                  f'{get_column_letter(column)}: original width differs')
    for row in range(1, original.max_row + 1):
        expected_height, actual_height = height(original, row), height(actual, row)
        if expected_height is not None:
            check(actual_height is not None and abs(expected_height - actual_height) < .01,
                  f'row {row}: original height differs')

    # All header cells, including blank spacer columns, and original fixed row
    # labels are checked at their source coordinates. Observations are excluded.
    structural = {(row, column) for row in range(1, spec['header'] + 1)
                  for column in range(1, original.max_column + 1)}
    structural.update((row, column) for row in range(spec['header'] + 1, spec['total'] + 1)
                      for column in (1, 2))
    for row, column in sorted(structural):
        ref = f'{get_column_letter(column)}{row}'
        expected = expected_label(source, cached, name, ref)
        received = actual[ref].value
        if ref in spec['dynamic']:
            # Exact date behavior is also exercised with different periods by
            # the Node integration test. The default selected year is 2026.
            if ref in ('M2', 'D3', 'D4'):
                check(received is not None and '3' in str(received), f'{ref}: selected quarter missing')
            else:
                check(received is not None and '2026' in str(received), f'{ref}: selected year missing')
        else:
            check(canonical(expected) == canonical(received), f'{ref}: original static cell differs')
        if expected is not None:
            expected_style, actual_style = style(original[ref]), style(actual[ref])
            changed = [key for key, value in expected_style.items() if value != actual_style[key]]
            check(not changed, f'{ref}: original style differs ({", ".join(changed)})')

    for row in actual:
        for cell in row:
            if cell.value is None:
                continue
            check(cell.column <= spec['width'], f'{cell.coordinate}: extra meaningful field')
            check(cell.font.name == 'Times New Roman', f'{cell.coordinate}: font differs')
            check(cell.data_type not in ('f', 'e'), f'{cell.coordinate}: source formula or Excel error exported')
            if cell.row > spec['total']:
                check(False, f'{cell.coordinate}: source observation outside the report body')
    for field in PAGE_FIELDS:
        check(getattr(actual.page_setup, field) == getattr(original.page_setup, field),
              f'pageSetup.{field}: original setting differs')
    for field in MARGIN_FIELDS:
        expected, received = getattr(original.page_margins, field), getattr(actual.page_margins, field)
        check(expected == received or (expected is not None and received is not None and abs(expected - received) < 1e-8),
              f'pageMargins.{field}: original margin differs')
    check(str(actual.print_area) == str(original.print_area), 'original print area differs')
    check(actual.print_title_rows == original.print_title_rows, 'original print title rows differ')
    check(actual.print_title_cols == original.print_title_cols, 'original print title columns differ')
    return {'sheet': name, 'failures': list(dict.fromkeys(failures))}


def numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def compare_summary_links(source, exported):
    """Independently follow the original source's cross-sheet total references."""
    failures = []
    for row in source['ИТОГО'].iter_rows(min_row=4, max_row=6, min_col=3, max_col=44):
        for cell in row:
            if cell.data_type != 'f':
                continue
            match = re.fullmatch(r"='?([^'!]+)'?!\$?([A-Z]+)\$?(\d+)", str(cell.value))
            if not match:
                continue
            source_name, column, source_row = match.groups()
            expected = exported[source_name][f'{column}{source_row}'].value
            actual = exported['ИТОГО'][cell.coordinate].value
            equivalent = abs(actual - expected) < 1e-8 if numeric(actual) and numeric(expected) else actual == expected
            if not equivalent:
                failures.append(f'{cell.coordinate}: cross-sheet total differs')
    return failures


def privacy_scan(source, exported, repo):
    # The workbook has three operational notes. Compare privately in memory;
    # report counts only, never the text or an identifying substring.
    notes = [str(source[name][ref].value) for name, refs in PRIVATE_NOTES.items()
             for ref in refs if source[name][ref].value is not None]
    allowed_extensions = {'.js', '.mjs', '.py', '.md', '.html', '.css', '.json'}
    files = [path for path in repo.rglob('*') if path.suffix in allowed_extensions
             and not any(part.startswith('.') or part == 'node_modules' for part in path.relative_to(repo).parts)
             and path.is_file()]
    contents = [path.read_text(errors='replace') for path in files]
    exported_values = [str(cell.value) for sheet in exported for row in sheet for cell in row if cell.value is not None]
    matches = sum(any(note in text for text in contents + exported_values) for note in notes)
    return {'sourceNotesChecked': len(notes), 'repositoryFilesChecked': len(files), 'matches': matches}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--no-generate', action='store_true')
    args = parser.parse_args()
    repo = Path(__file__).resolve().parents[1]
    if not args.no_generate:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        code = """
import {writeFile} from 'node:fs/promises';
import {generateArchiveReport} from './archive-engine.js';
import {exportArchiveXlsx} from './export.js';
await writeFile(process.argv[1], exportArchiveXlsx(generateArchiveReport('cooperation-sva-summary')));
"""
        subprocess.run([os.environ.get('CODEX_PRIMARY_RUNTIME_NODE', 'node'), '--input-type=module', '-e', code,
                        str(args.output.resolve())], cwd=repo, check=True)
    source = openpyxl.load_workbook(args.source, data_only=False)
    cached = openpyxl.load_workbook(args.source, data_only=True)
    exported = openpyxl.load_workbook(args.output, data_only=False)
    workbook_failures = []
    if exported.sheetnames != source.sheetnames:
        workbook_failures.append('Original worksheet names/order/count differ')
    if set(exported.sheetnames) != set(SHEETS):
        workbook_failures.append('Missing source sheet or extra technical worksheet')
    results = [compare_sheet(source, cached, exported, name) for name in SHEETS if name in exported.sheetnames]
    if all(name in exported.sheetnames for name in SHEETS):
        workbook_failures.extend(compare_summary_links(source, exported))
    privacy = privacy_scan(source, exported, repo)
    if privacy['matches']:
        workbook_failures.append('Private source notes found in the repository or export')
    report = {'passed': not workbook_failures and all(not result['failures'] for result in results),
              'workbookFailures': workbook_failures, 'sheets': results, 'privacy': privacy}
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report['passed'] else 1


if __name__ == '__main__':
    raise SystemExit(main())
