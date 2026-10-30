#!/usr/bin/env python3
"""Extract the four-sheet SVA source layout, never customer financial facts.

Usage: python scripts/extract-cooperation-form.py '/path/to/Итоговая 4 ЦГО, МИО и ДП .xlsx'

Allowlist: structural headers, institution names, ordinal numbers and total
captions. Two institution labels are read from string caches at explicit cells.
No numeric formula cache, external formula, observation, footer signature,
comment or hyperlink is included. The source-layout helper forces TNR.
"""
from pathlib import Path
import argparse
import importlib.util
import json
import re
import openpyxl
from openpyxl.utils.cell import get_column_letter, range_boundaries

SHEETS = [
    ('ЦГО', 'cgo', 3, 4, 33, 34),
    ('МИО', 'mio', 3, 4, 23, 24),
    ('ДП', 'dp', 5, 6, 25, 26),
    ('ИТОГО', 'summary', 3, 4, 6, 7),
]
FORM_ID = 'cooperation-sva-summary'
TITLE = 'Итоги деятельности СВА (приложение 1)'


def clean(value):
    return re.sub(r'\s+', ' ', str(value)).strip()


def headers(ws, values, depth):
    cells = {ref: value for ref, value in values.items() if int(re.search(r'\d+$', ref)[0]) <= depth}
    merges = []
    for ref in ws.merged_cells.ranges:
        c1, r1, c2, r2 = range_boundaries(str(ref))
        if r2 <= depth:
            merges.append(str(ref))
    columns = []
    # Exclude the document caption/period row from flattened column labels.
    first_label_row = 4 if ws.title == 'ДП' else 2
    for col in range(1, ws.max_column + 1):
        labels, source_cells = [], []
        for row in range(first_label_row, depth + 1):
            ref = f'{get_column_letter(col)}{row}'
            label = cells.get(ref)
            if label is None:
                for merged in merges:
                    c1, r1, c2, r2 = range_boundaries(merged)
                    if c1 <= col <= c2 and r1 <= row <= r2:
                        ref = f'{get_column_letter(c1)}{r1}'
                        label = cells.get(ref)
                        break
            if isinstance(label, str) and clean(label):
                text = clean(label)
                if text not in labels:
                    labels.append(text)
                if ref not in source_cells:
                    source_cells.append(ref)
        columns.append({'letter': get_column_letter(col), 'labels': labels, 'sourceCells': source_cells})
    return {'columnCount': ws.max_column, 'headerDepth': depth, 'sourceHeaderRows': [1, depth],
            'headerCells': cells, 'merges': sorted(merges), 'columns': columns}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('source', type=Path)
    parser.add_argument('--output', type=Path, default=Path('cooperation-schemas.js'))
    args = parser.parse_args()
    spec = importlib.util.spec_from_file_location('source_layout', Path(__file__).with_name('source-layout.py'))
    helper = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(helper)
    source = openpyxl.load_workbook(args.source, data_only=False)
    cached = openpyxl.load_workbook(args.source, data_only=True)
    assert source.sheetnames == [sheet[0] for sheet in SHEETS], 'Source sheet order changed; review the allowlist.'
    sheets = []
    for sheet_name, key, depth, start, end, total in SHEETS:
        ws = source[sheet_name]
        allowed = {cell.coordinate for row in ws.iter_rows(min_row=1, max_row=depth) for cell in row
                   if cell.value is not None and cell.data_type != 'f'}
        allowed.update(f'{letter}{row}' for row in range(start, end + 1) for letter in ['A', 'B'])
        allowed.add(f'A{total}')
        layout = helper.extract_source_layout(args.source, sheet_name=sheet_name, allowed_cells=allowed)
        resolved_cells = []
        # These are names of institutions, never amounts or measurements.
        if sheet_name == 'ЦГО':
            for ref in ['B31', 'B32']:
                label = cached[sheet_name][ref].value
                if not isinstance(label, str) or not label.strip() or label.startswith(('=', '#')):
                    raise ValueError(f'Missing institution label cache: {sheet_name}!{ref}')
                layout['cells'][ref] = label
                resolved_cells.append({'cell': ref, 'kind': 'cached-institution-label'})
        if sheet_name == 'МИО':
            # This source formula is a local reference to a structural header.
            assert ws['AU2'].value == '=ЦГО!$AQ$2', 'Certificate header reference changed.'
            label = source['ЦГО']['AQ2'].value
            assert isinstance(label, str) and label.strip()
            layout['cells']['AU2'] = label
            resolved_cells.append({'cell': 'AU2', 'kind': 'local-header-reference', 'sourceSheet': 'ЦГО', 'sourceCell': 'AQ2'})
        data_rows = []
        for row in range(start, end + 1):
            ordinal = layout['cells'].get(f'A{row}')
            label = layout['cells'].get(f'B{row}')
            if not isinstance(ordinal, (int, float)) or not isinstance(label, str):
                raise ValueError(f'Unexpected institution identity in {sheet_name}!A{row}:B{row}')
            data_rows.append({'id': f'{key}-{int(ordinal):02d}', 'sourceRow': row, 'ordinal': int(ordinal),
                              'label': label, 'labelCell': f'B{row}', 'ordinalCell': f'A{row}'})
        schema = {'id': f'{FORM_ID}-{key}', 'key': key, 'category': sheet_name,
                  'name': sheet_name, 'sourceSheetName': sheet_name,
                  'kind': 'overall-summary' if key == 'summary' else 'institution-summary',
                  **headers(ws, layout['cells'], depth),
                  'dataRows': data_rows,
                  'totalRows': [{'id': f'{key}-total', 'sourceRow': total, 'label': layout['cells'][f'A{total}'],
                                 'labelCell': f'A{total}', 'fromRow': start, 'toRow': end}],
                  'templateLayout': layout}
        schema['valueColumns'] = [column['letter'] for column in schema['columns'][2:] if column['labels']]
        if resolved_cells:
            schema['resolvedLabelCells'] = resolved_cells
        if key == 'summary':
            links = []
            for row in ws.iter_rows(min_row=start, max_row=end, min_col=3):
                for cell in row:
                    if cell.data_type != 'f':
                        continue
                    # Only literal same-workbook single-cell references are
                    # reported as structural metadata. They are not executable.
                    match = re.fullmatch(r"=('?)(ЦГО|МИО|ДП)\1!\$?([A-Z]+)\$?(\d+)", cell.value)
                    if match:
                        links.append({'target': cell.coordinate, 'sourceSheet': match[2],
                                      'sourceCell': f'{match[3]}{match[4]}'})
            schema['summaryLinks'] = links
        sheets.append(schema)
    form = {'id': FORM_ID, 'name': TITLE, 'sourceFile': args.source.name,
            'departmentId': 'audit-cooperation', 'catalogNumber': 1, 'sources': ['СВА'],
            'engine': 'cooperation', 'sheets': sheets}
    args.output.write_text('// Source structure only. Institution labels and headers are retained; customer observations, external formulas and signatures are excluded.\nexport const cooperationForms = '
                          + json.dumps([form], ensure_ascii=False, separators=(',', ':')) + ';\n')
    print(json.dumps({'form': FORM_ID, 'sheets': [{'name': sheet['name'], 'rows': sheet['templateLayout']['rowCount'],
                      'columns': sheet['columnCount'], 'institutions': len(sheet['dataRows']),
                      'totalRow': sheet['totalRows'][0]['sourceRow']} for sheet in sheets]}, ensure_ascii=False))


if __name__ == '__main__':
    main()
