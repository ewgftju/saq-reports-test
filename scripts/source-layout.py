"""Read a sanitized Excel layout. No customer observations are copied by default.

extract_source_layout(path, sheet_name=None, column_count=None, row_count=None,
                      allowed_cells=()) -> templateLayout

Only explicit allowed_cells can contribute literal values. All cells may
contribute formatting; formulas, hyperlinks, comments and document properties
are never carried into the result. The user requested Times New Roman globally.
"""
from pathlib import Path
import struct
import xml.etree.ElementTree as ET
import openpyxl
import xlrd
from openpyxl.utils.cell import get_column_letter, range_boundaries


def _compact(values):
    return {key: value for key, value in values.items() if value is not None}


def _rgb(rgb):
    if rgb is None:
        return None
    text = str(rgb).replace('#', '')
    return 'FF' + text[-6:] if len(text) in (6, 8) else None


def _theme_colors(workbook):
    if not workbook.loaded_theme:
        return []
    root = ET.fromstring(workbook.loaded_theme)
    scheme = root.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}clrScheme')
    if scheme is None:
        return []
    # Spreadsheet theme indexing starts light1/dark1, unlike XML order.
    names = ['lt1', 'dk1', 'lt2', 'dk2', 'accent1', 'accent2', 'accent3', 'accent4', 'accent5', 'accent6', 'hlink', 'folHlink']
    result = []
    for name in names:
        parent = scheme.find('{http://schemas.openxmlformats.org/drawingml/2006/main}' + name)
        child = list(parent)[0] if parent is not None and len(parent) else None
        result.append(_rgb(child.get('val') if child is not None and child.tag.endswith('srgbClr') else child.get('lastClr') if child is not None else None))
    return result


def _xlsx_color(value, themes):
    if value is None:
        return None
    if value.type == 'rgb':
        result = _rgb(value.rgb)
    elif value.type == 'indexed':
        result = _rgb(openpyxl.styles.colors.COLOR_INDEX[value.indexed]) if value.indexed < len(openpyxl.styles.colors.COLOR_INDEX) else None
    elif value.type == 'theme':
        result = themes[value.theme] if value.theme < len(themes) else None
    else:
        result = None
    if result and value.tint:
        channels = [int(result[index:index + 2], 16) for index in (2, 4, 6)]
        tint = value.tint
        channels = [round(channel * (1 + tint) if tint < 0 else channel + (255 - channel) * tint) for channel in channels]
        return 'FF' + ''.join(f'{max(0, min(255, channel)):02X}' for channel in channels)
    return result


def _xlsx_style(cell, themes):
    font = cell.font
    fill = cell.fill
    alignment = cell.alignment
    border = {}
    for edge in ['left', 'right', 'top', 'bottom']:
        side = getattr(cell.border, edge)
        if side and side.style:
            border[edge] = _compact({'style': side.style, 'color': _xlsx_color(side.color, themes)})
    return _compact({
        'font': _compact({'name': 'Times New Roman', 'size': font.sz or 11, 'bold': bool(font.b), 'italic': bool(font.i), 'color': _xlsx_color(font.color, themes), 'underline': font.u, 'strike': bool(font.strike)}),
        'fill': {'color': _xlsx_color(fill.fgColor, themes)} if fill.patternType == 'solid' else None,
        'border': border,
        'alignment': _compact({'horizontal': alignment.horizontal, 'vertical': alignment.vertical, 'wrapText': bool(alignment.wrap_text), 'textRotation': alignment.text_rotation, 'shrinkToFit': bool(alignment.shrink_to_fit), 'indent': alignment.indent}),
        'numberFormat': cell.number_format or 'General'
    })


def _literal(value):
    if value is None or isinstance(value, str) and value.startswith('='):
        return None
    if isinstance(value, (str, int, float, bool)):
        return value
    return None


def _xlsx(path, sheet_name, column_count, row_count, allowed):
    workbook = openpyxl.load_workbook(path, data_only=False)
    sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
    columns = column_count or sheet.max_column
    rows = row_count or sheet.max_row
    themes = _theme_colors(workbook)
    styles, style_keys, cell_styles, cells = [], {}, {}, {}
    for row in sheet.iter_rows(min_row=1, max_row=rows, min_col=1, max_col=columns):
        for cell in row:
            style = _xlsx_style(cell, themes)
            key = repr(style)
            if key not in style_keys:
                style_keys[key] = len(styles)
                styles.append(style)
            cell_styles[cell.coordinate] = style_keys[key]
            if cell.coordinate in allowed:
                value = _literal(cell.value)
                if value is not None:
                    cells[cell.coordinate] = value
    merges = [str(ref) for ref in sheet.merged_cells.ranges if range_boundaries(str(ref))[2] <= columns and range_boundaries(str(ref))[3] <= rows]
    widths, hidden_columns = {}, []
    for column in range(1, columns + 1):
        letter = get_column_letter(column)
        dimensions = [dimension for dimension in sheet.column_dimensions.values() if dimension.min and dimension.min <= column <= (dimension.max or dimension.min)]
        dimension = dimensions[0] if dimensions else None
        widths[letter] = dimension.width if dimension is not None else sheet.sheet_format.defaultColWidth or 8.43
        if dimension is not None and dimension.hidden:
            hidden_columns.append(letter)
    heights = {str(row): sheet.row_dimensions[row].height or sheet.sheet_format.defaultRowHeight or 15 for row in range(1, rows + 1)}
    setup = _compact({key: getattr(sheet.page_setup, key) for key in ['orientation', 'paperSize', 'scale', 'fitToWidth', 'fitToHeight', 'firstPageNumber', 'useFirstPageNumber', 'blackAndWhite', 'draft', 'horizontalDpi', 'verticalDpi', 'pageOrder']})
    page_properties = sheet.sheet_properties.pageSetUpPr
    return {'columnCount': columns, 'rowCount': rows, 'cells': cells, 'cellStyles': cell_styles, 'styles': styles, 'merges': sorted(merges), 'columnWidths': widths, 'rowHeights': heights, 'defaultRowHeight': sheet.sheet_format.defaultRowHeight or 15, 'defaultColWidth': sheet.sheet_format.defaultColWidth or 8.43, 'hiddenColumns': hidden_columns, 'hiddenRows': [row for row in range(1, rows + 1) if sheet.row_dimensions[row].hidden], 'sheetView': _compact({key: getattr(sheet.sheet_view, key) for key in ['showGridLines', 'showRowColHeaders', 'zoomScale', 'zoomScaleNormal']}), 'pageSetup': setup, 'pageSetupProperties': _compact({key: getattr(page_properties, key) for key in ['fitToPage', 'autoPageBreaks']}) if page_properties is not None else {}, 'pageMargins': {key: getattr(sheet.page_margins, key) for key in ['left', 'right', 'top', 'bottom', 'header', 'footer']}, 'printOptions': _compact({key: getattr(sheet.print_options, key) for key in ['horizontalCentered', 'verticalCentered', 'headings', 'gridLines']}), 'printArea': str(sheet.print_area) or None, 'printTitlesRows': sheet.print_title_rows, 'printTitlesCols': sheet.print_title_cols}


def _xls_color(book, index):
    value = book.colour_map.get(index)
    return 'FF' + ''.join(f'{channel:02X}' for channel in value) if value is not None else None


def _xls_style(book, xf):
    font = book.font_list[xf.font_index]
    border = xf.border
    names = {1: 'thin', 2: 'medium', 3: 'dashed', 4: 'dotted', 5: 'thick', 6: 'double', 7: 'hair', 8: 'mediumDashed', 9: 'dashDot', 10: 'mediumDashDot', 11: 'dashDotDot', 12: 'mediumDashDotDot', 13: 'slantDashDot'}
    sides = {}
    for edge in ['left', 'right', 'top', 'bottom']:
        line = getattr(border, edge + '_line_style')
        if line:
            sides[edge] = _compact({'style': names.get(line, 'thin'), 'color': _xls_color(book, getattr(border, edge + '_colour_index'))})
    align = xf.alignment
    return _compact({'font': _compact({'name': 'Times New Roman', 'size': font.height / 20, 'bold': bool(font.bold), 'italic': bool(font.italic), 'color': _xls_color(book, font.colour_index), 'underline': {1: 'single', 2: 'double', 33: 'singleAccounting', 34: 'doubleAccounting'}.get(font.underline_type), 'strike': bool(font.struck_out)}), 'fill': {'color': _xls_color(book, xf.background.pattern_colour_index)} if xf.background.fill_pattern == 1 else None, 'border': sides, 'alignment': _compact({'horizontal': {0: 'general', 1: 'left', 2: 'center', 3: 'right', 4: 'fill', 5: 'justify', 6: 'centerContinuous', 7: 'distributed'}.get(align.hor_align), 'vertical': {0: 'top', 1: 'center', 2: 'bottom', 3: 'justify', 4: 'distributed'}.get(align.vert_align), 'wrapText': bool(align.text_wrapped), 'textRotation': align.rotation, 'shrinkToFit': bool(align.shrink_to_fit), 'indent': align.indent_level}), 'numberFormat': book.format_map[xf.format_key].format_str if xf.format_key in book.format_map else 'General'})


def _xls_print(path, book, index):
    from xlrd.compdoc import CompDoc
    compound = CompDoc(Path(path).read_bytes())
    stream = compound.get_named_stream('Workbook') or compound.get_named_stream('Book')
    position = book._sh_abs_posn[index] - book.base
    margins, setup, options = {}, {}, {}
    while position + 4 <= len(stream):
        opcode, length = struct.unpack_from('<HH', stream, position)
        data = stream[position + 4:position + 4 + length]
        position += 4 + length
        if opcode == 0x000A:
            break
        if opcode in {0x26, 0x27, 0x28, 0x29} and len(data) >= 8:
            margins[{0x26: 'left', 0x27: 'right', 0x28: 'top', 0x29: 'bottom'}[opcode]] = struct.unpack_from('<d', data)[0]
        if opcode in {0x2A, 0x2B, 0x83, 0x84} and len(data) >= 2:
            options[{0x2A: 'headings', 0x2B: 'gridLines', 0x83: 'horizontalCentered', 0x84: 'verticalCentered'}[opcode]] = bool(struct.unpack_from('<H', data)[0])
        if opcode == 0xA1 and len(data) >= 34:
            paper, scale, page, width, height, flags, hres, vres, header, footer, copies = struct.unpack_from('<8H2dH', data)
            setup = {'paperSize': paper, 'scale': scale, 'firstPageNumber': page, 'fitToWidth': width, 'fitToHeight': height, 'orientation': 'portrait' if flags & 2 else 'landscape', 'blackAndWhite': bool(flags & 8), 'draft': bool(flags & 16), 'useFirstPageNumber': bool(flags & 128), 'pageOrder': 'overThenDown' if flags & 1 else 'downThenOver', 'horizontalDpi': hres, 'verticalDpi': vres}
            margins.update(header=header, footer=footer)
    return setup, margins, options


def _xls(path, sheet_name, column_count, row_count, allowed):
    book = xlrd.open_workbook(path, formatting_info=True)
    sheet = book.sheet_by_name(sheet_name) if sheet_name else book.sheet_by_index(0)
    index = book.sheet_names().index(sheet.name)
    columns, rows = column_count or sheet.ncols, row_count or sheet.nrows
    styles, style_keys, cell_styles, cells = [], {}, {}, {}
    for row in range(rows):
        for column in range(columns):
            ref = f'{get_column_letter(column + 1)}{row + 1}'
            xf_index = sheet.cell_xf_index(row, column) if row < sheet.nrows and column < sheet.ncols else 0
            style = _xls_style(book, book.xf_list[xf_index])
            key = repr(style)
            if key not in style_keys:
                style_keys[key] = len(styles)
                styles.append(style)
            cell_styles[ref] = style_keys[key]
            if ref in allowed and row < sheet.nrows and column < sheet.ncols:
                value = _literal(sheet.cell_value(row, column))
                if value is not None and value != '':
                    cells[ref] = value
    merges = [f'{get_column_letter(c + 1)}{r + 1}:{get_column_letter(cc)}{rr}' for r, rr, c, cc in sheet.merged_cells if rr <= rows and cc <= columns]
    setup, margins, options = _xls_print(path, book, index)
    widths = {get_column_letter(col + 1): sheet.colinfo_map[col].width / 256 if col in sheet.colinfo_map else sheet.defcolwidth or 8.43 for col in range(columns)}
    heights = {str(row + 1): sheet.rowinfo_map[row].height / 20 if row in sheet.rowinfo_map else sheet.default_row_height / 20 for row in range(rows)}
    return {'columnCount': columns, 'rowCount': rows, 'cells': cells, 'cellStyles': cell_styles, 'styles': styles, 'merges': merges, 'columnWidths': widths, 'rowHeights': heights, 'hiddenColumns': [get_column_letter(col + 1) for col in range(columns) if col in sheet.colinfo_map and sheet.colinfo_map[col].hidden], 'hiddenRows': [row + 1 for row in range(rows) if row in sheet.rowinfo_map and sheet.rowinfo_map[row].hidden], 'pageSetup': setup, 'pageMargins': margins, 'printOptions': options, 'printArea': None, 'printTitlesRows': None, 'printTitlesCols': None}


def extract_source_layout(path, sheet_name=None, column_count=None, row_count=None, allowed_cells=()):
    path = Path(path)
    reader = _xls if path.suffix.lower() == '.xls' else _xlsx
    return reader(path, sheet_name, column_count, row_count, set(allowed_cells))
