#!/usr/bin/env python3
"""Extract structural schemas only; never historical customer observations.

Usage: python scripts/extract-archive-forms.py /path/to/УАО.zip [archive-schemas.js]
The per-workbook allowlist is deliberate: text below headers can contain PII.
openpyxl/xlrd are used only to READ workbooks. No workbook is authored here.
"""
from __future__ import annotations
import argparse
import io
import json
import re
import warnings
from pathlib import Path
from zipfile import ZipFile
import openpyxl
import xlrd
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter, range_boundaries

# filename suffix, stable id, catalog number, title, family, header rows, kind
MANIFEST = [
 ('15. Форма критериев.xlsx','uao-rating-criteria',15,'Рейтинг — форма критериев','rating',None,'rating'),
 ('16.Информация по исполнению опер плана.xls','uao-operplan',16,'Исполнение операционного плана','operplan',(3,5),'table'),
 ('17. Данные по итогам исполнения рб.xlsx','uao-budget-execution',17,'Итоги аудита исполнения республиканского бюджета — приложение 26','budget-execution',(11,12),'matrix'),
 ('18. Фин нарушения приципов.xlsx','uao-budget-principles',18,'Финансовые нарушения принципов бюджетной системы — приложение 27','budget-principles',(11,13),'table'),
 ('19. информация о нарушениях, выявленных в гос уч.xlsx','uao-budget-violations',19,'Нарушения при использовании средств республиканского бюджета — приложение 72','budget-violations',(11,13),'table'),
 ('21. Форма по остаткам/форма остатка.xlsx','uao-residuals',21,'Остатки нарушений — общая форма','residuals',(3,5),'table'),
 ('21. Форма по остаткам/форма по остатку 2.xlsx','uao-residuals-years',21,'Расшифровка остатков невозмещенных сумм по годам','residuals-years',(4,10),'table'),
 ('22. СОЦ эконом/Таблица.xls','uao-social',22,'Итоги деятельности КВГА — Соцэконом','social',(5,5),'matrix'),
 ('6 Итоги аудиторской деятельности ОВГА в разрезе  дел АБП.xlsx','uao-abp',6,'Итоги аудиторской деятельности в разрезе дел АБП','audit-abp',(4,9),'table'),
 ('7 Итоги аудиторской деятельности ОВГА в разрезе  дел ВГА.xlsx','uao-vga',7,'Итоги аудиторской деятельности в разрезе дел ВГА','audit-cases',(4,9),'table'),
 ('8 Итоги аудиторской деятельности ОВГА в разрезе  дел органов аудита.xlsx','uao-org',8,'Итоги аудиторской деятельности в разрезе органов аудита','audit-orgs',(4,9),'table'),
 ('дисциплинарка 4.xlsx','uao-discipline-persons',4,'Лица, привлеченные к дисциплинарной ответственности','discipline-persons',(5,5),'table'),
 ('Оперплан/отчетность по Опер плану нарушения.xls','uao-operplan-violations',16,'Операционный план — восстановление и возмещение нарушений','operplan-violations',(3,4),'table'),
 ('Оперплан/отчетность по Опер плану предписания.xls','uao-operplan-orders',16,'Операционный план — исполнение предписаний','operplan-orders',(3,4),'table'),
 ('Отчёт_по_видам_нарушений.xlsx','uao-violation-types',10,'Отчет по видам нарушений','violation-types',(2,6),'table'),
 ('Отчет по основаниям.xls','uao-grounds',12,'Отчет по основаниям проведения аудита','grounds',(3,5),'table'),
 ('Отчет_по_государственным_закупкам.xlsx','uao-procurement',11,'Отчет по государственным закупкам','procurement',(4,10),'table'),
 ('Отчет_по_Прочим_нарушениям.xlsx','uao-other-violations',13,'Отчет по прочим нарушениям','other-violations',(4,8),'table'),
 ('рейтинг адм дисциплинарка/4,9  Дисциплинарка адм/Приложение 2 по Дисцип0731173317.xlsx','uao-discipline-summary',4,'Дисциплинарные взыскания — свод по видам','discipline-summary',(4,7),'table'),
 ('рейтинг адм дисциплинарка/4,9  Дисциплинарка адм/приложение 3 АП0731173356.xlsx','uao-administrative',9,'Информация об административном правонарушении','administrative',(4,6),'table'),
 ('рейтинг адм дисциплинарка/Рейтинг/Рейтинг измен 20250528154137.xlsx','uao-rating-revised',15,'Рейтинг — измененная форма из архива','rating',None,'rating'),
 ('Функциональный_анализ_по_аудиту.xlsx','uao-functional',14,'Функциональный анализ по аудиту','functional',(2,6),'table'),
 ('ЭГА.xlsx','uao-ega',20,'Итоги электронного государственного аудита','electronic-audit',(4,9),'table'),
]

def clean(value):
    return re.sub(r'\s+', ' ', str(value).replace('\u200b', '')).strip()

def position(ref):
    col, row = coordinate_from_string(ref)
    return column_index_from_string(col), row

def decoded_name(info):
    return info.filename if info.flag_bits & 0x800 else info.filename.encode('cp437').decode('cp866')

def read_workbook(content, is_xlsx):
    result = []
    if is_xlsx:
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=False, read_only=False)
        for sheet in workbook:
            cells = {c.coordinate:c.value for row in sheet for c in row if c.value is not None}
            result.append({'name':sheet.title,'rows':sheet.max_row,'cols':sheet.max_column,'cells':cells,'merges':[str(m) for m in sheet.merged_cells.ranges]})
    else:
        workbook = xlrd.open_workbook(file_contents=content, formatting_info=True)
        for sheet in workbook.sheets():
            cells={f'{get_column_letter(c+1)}{r+1}':sheet.cell_value(r,c) for r in range(sheet.nrows) for c in range(sheet.ncols) if sheet.cell_value(r,c)!=''}
            merges=[f'{get_column_letter(c+1)}{r+1}:{get_column_letter(cc)}{rr}' for r,rr,c,cc in sheet.merged_cells]
            result.append({'name':sheet.name,'rows':sheet.nrows,'cols':sheet.ncols,'cells':cells,'merges':merges})
    return result

def table_header(source, first, last, family):
    raw=source['cells']; header={}; merges=[]
    for ref,value in raw.items():
        col,row=position(ref)
        if first<=row<=last and isinstance(value,str) and not value.startswith('='):
            label=clean(value)
            if not label:continue
            if family=='operplan' and re.search(r'Факт на',label,re.I):label='Факт за выбранный период'
            if family=='budget-execution' and re.search(r'\b20\d\d год',label):label='Отчетный период, всего'
            if family=='social' and re.search(r'на 1 октября',label):label='Отчетный период' if col==6 else 'Предыдущий период'
            header[f'{get_column_letter(col)}{row-first+1}']=label
    for ref in source['merges']:
        c1,r1,c2,r2=range_boundaries(ref)
        if first<=r1<=r2<=last:merges.append(f'{get_column_letter(c1)}{r1-first+1}:{get_column_letter(c2)}{r2-first+1}')
    if family=='grounds':header.setdefault('C2','Орган аудита')
    if family=='rating' and source['name']=='показатели':header.setdefault('C1','№ строки')
    last_col=max([position(ref)[0] for ref in header]+[range_boundaries(ref)[2] for ref in merges])
    columns=[]
    for col in range(1,last_col+1):
        labels=[]
        for row in range(1,last-first+2):
            ref=f'{get_column_letter(col)}{row}';value=header.get(ref)
            if value is None:
                for merge in merges:
                    c1,r1,c2,r2=range_boundaries(merge)
                    if c1<=col<=c2 and r1<=row<=r2:
                        value=header.get(f'{get_column_letter(c1)}{r1}');break
            if value and value not in labels:labels.append(value)
        columns.append({'letter':get_column_letter(col),'labels':labels})
    return {'columnCount':last_col,'sourceColumnCount':source['cols'],'headerDepth':last-first+1,'headerCells':header,'merges':merges,'columns':columns,'sourceHeaderRows':[first,last]}

def matrix_rows(source, family):
    cells=source['cells'];sheet=source['name'];rows=[]
    if family=='budget-execution':start,end,static_cols,values=14,30,['A'],['B','C','D','E']
    elif family=='social':start,end,static_cols,values=6,65,list('ABCDE'),list('FGHI')
    elif sheet=='показатели':start,end,static_cols,values=5,143,list('ABCDE'),[get_column_letter(c) for c in range(6,26)]
    elif sheet=='баллы':start,end,static_cols,values=4,384,list('ABCD'),[get_column_letter(c) for c in range(5,25)]
    elif sheet=='итоговые баллы':start,end,static_cols,values=4,73,list('ABCD'),[get_column_letter(c) for c in range(5,25)]
    else:raise ValueError(f'Unexpected matrix {family}/{sheet}')
    active_section='';active_code=''
    for n in range(start,end+1):
        if family=='social' and n in [38,39,40]:continue
        static={}
        for col in static_cols:
            value=cells.get(f'{col}{n}')
            if value is None or (isinstance(value,str) and value.startswith('=')):continue
            if family=='rating' and col=='C' and sheet in ['баллы','итоговые баллы']:continue
            # The numeric cells on this allowlist are ordinal/criterion codes.
            if isinstance(value,(int,float)) and (family!='rating' or col not in ['A','B','C']):continue
            text=clean(int(value) if isinstance(value,float) and value.is_integer() else value)
            if text and text!=',':static[col]=text
        if not static:continue
        preferred='A' if family=='budget-execution' else 'D' if sheet=='показатели' else 'B'
        label=static.get(preferred) or static.get('B') or static.get('A') or ''
        if family=='social':label=' / '.join(v for c,v in static.items() if c in ['B','C','D']) or static.get('A','')
        section=bool(re.search(r'^(?:Направление\s+\d|\d+-направление)',label,re.I))
        role='section' if section else 'subtotal' if re.search(r'^Итоговые баллы$',label,re.I) else 'indicator'
        if section:active_section=label
        code_col='B' if sheet=='показатели' else 'A'
        code=static.get(code_col,'') if family=='rating' else ''
        if re.fullmatch(r'\d+\.\d+',code):active_code=code
        row_values=[] if section else values[:]
        if family=='rating' and sheet=='баллы':
            has_result=any(f'{col}{n}' in cells for col in values)
            row_values=values[:] if has_result and not section else []
            if not code and not section:role='scale'
        row={'id':f'r{n}','label':label,'sourceRow':n,'cells':static,'valueColumns':row_values,'role':role}
        if active_section:row['section']=active_section
        if active_code and family=='rating':row['code']=active_code
        if family=='social' and static.get('E'):row['unit']=static['E']
        if family=='rating' and sheet in ['баллы','итоговые баллы']:
            points=cells.get(f'C{n}')
            if points is not None and not (isinstance(points,str) and points.startswith('=')):
                row['rulePoints']={'value':points,'sourceCell':f'C{n}','column':'C'}
        rows.append(row)
    # Add vertically merged parent labels to the human-readable mapping context.
    for row in rows:
        if family!='social':continue
        n=row['sourceRow'];parts=[]
        for col in ['B','C','D']:
            value=cells.get(f'{col}{n}')
            if value is None:
                ci=column_index_from_string(col)
                for merged in source['merges']:
                    c1,r1,c2,r2=range_boundaries(merged)
                    if c1<=ci<=c2 and r1<=n<=r2:value=cells.get(f'{get_column_letter(c1)}{r1}');break
            if isinstance(value,str) and not value.startswith('=') and clean(value) not in parts:parts.append(clean(value))
        if parts:row['label']=' / '.join(parts)
    return rows

def make_form(source_file, definition, sources):
    suffix,id,number,name,family,header,kind=definition
    form={'id':id,'sourceFile':source_file,'uaoNumber':number,'name':name,'family':family,'sheets':[]}
    for i,source in enumerate(sources):
        source_name=source['name'];current_kind=kind;current_header=header
        if kind=='rating':
            current_header=(3,4) if source_name=='показатели' else (3,3) if source_name=='итоги' else (2,3)
            current_kind='table' if source_name=='итоги' else 'matrix'
        layout=table_header(source,*current_header,family)
        normalized_name=source_name.strip()
        if family=='operplan':normalized_name='Операционный план'
        elif family=='social':normalized_name='Соцэконом'
        elif family=='residuals':normalized_name='Остатки нарушений'
        schema={'id':f'{id}-s{i+1}','name':normalized_name,'sourceSheetName':source_name,'kind':current_kind,**layout}
        if current_kind=='matrix':schema['rows']=matrix_rows(source,family)
        if family in ['budget-principles','budget-violations']:
            schema['rowGroups']=[{'id':'state','label':'Государственные учреждения'},{'id':'quasi','label':'Субъекты квазигосударственного сектора'}] if family=='budget-principles' else [{'id':'state','label':'Государственные учреждения'},{'id':'local','label':'Местные исполнительные органы'},{'id':'quasi','label':'Субъекты квазигосударственного сектора'}]
        if family=='rating':
            if source_name=='итоги':schema['ratingBands']=['Превосходно','Эффективно','Удовлетворительно','Не эффективно']
            formula_count=sum(isinstance(v,str) and v.startswith('=') for v in source['cells'].values())
            invalid=sum(isinstance(v,str) and v.startswith('=') and ('#REF!' in v or '[' in v) for v in source['cells'].values())
            schema['formulaAudit']={'sourceFormulaCount':formula_count,'invalidOrExternalCount':invalid,'execution':'not-imported'}
        form['sheets'].append(schema)
    return form

def main():
    parser=argparse.ArgumentParser();parser.add_argument('archive');parser.add_argument('output',nargs='?',default='archive-schemas.js');args=parser.parse_args()
    forms=[]
    with ZipFile(args.archive) as archive:
        entries={decoded_name(info):info for info in archive.infolist() if decoded_name(info).lower().endswith(('.xls','.xlsx'))}
        expected={'УАО/'+item[0] for item in MANIFEST}
        if set(entries)!=expected:raise ValueError(f'Archive differs from reviewed allowlist: {set(entries)^expected}')
        for definition in MANIFEST:
            source_file='УАО/'+definition[0]
            forms.append(make_form(source_file,definition,read_workbook(archive.read(entries[source_file]),source_file.endswith('.xlsx'))))
    payload=json.dumps(forms,ensure_ascii=False,separators=(',',':'))
    assert not re.search(r'\b\d{12}\b',payload),'A customer identifier reached the structural schema'
    assert not re.search(r'\b\d{5}-\d{2}-\d{4,}\b',payload),'A historical audit case reached the structural schema'
    # Intern identical values and subtrees. Ratings and wide audit headers reuse
    # the same long labels; the lossless pool reduces the browser payload.
    pool=[];seen={}
    def pack(value):
        key=json.dumps(value,ensure_ascii=False,separators=(',',':'))
        if key in seen:return seen[key]
        if isinstance(value,list):node=[0]+[pack(v) for v in value]
        elif isinstance(value,dict):node=[1]+[part for k,v in value.items() for part in (pack(k),pack(v))]
        else:node=value
        index=len(pool);pool.append(node);seen[key]=index
        return index
    root=pack(forms)
    packed=json.dumps(pool,ensure_ascii=False,separators=(',',':'))
    decoder="const values=[];\nfor(const node of pool){\n if(!Array.isArray(node)){values.push(node);continue;}\n if(node[0]===0){values.push(node.slice(1).map(i=>values[i]));continue;}\n const value={};for(let i=1;i<node.length;i+=2)value[values[node[i]]]=values[node[i+1]];values.push(value);\n}\n"
    output=Path(args.output)
    output.write_text('// Generated by scripts/extract-archive-forms.py. Structural labels only; customer results are excluded.\nconst pool='+packed+';\n'+decoder+'export const archiveForms=values['+str(root)+'];\n',encoding='utf-8')
    print(json.dumps({'files':len(forms),'sheets':sum(len(f['sheets']) for f in forms),'output':str(output),'bytes':output.stat().st_size},ensure_ascii=False))

if __name__=='__main__':main()
